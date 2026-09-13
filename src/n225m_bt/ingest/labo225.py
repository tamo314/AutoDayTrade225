"""225Labo source adapter; it never fabricates a contract month or prices."""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Literal

import polars as pl

from n225m_bt.config import SourceFormat
from n225m_bt.ingest.source_detect import (
    REQUIRED,
    SourceInspection,
    inspect_delimited_bytes,
    inspection_from_headers,
)


@dataclass(frozen=True, slots=True)
class SourceRecord:
    source_file: str
    source_hash: str
    source_row_number: int
    trade_date: date
    source_time: time
    open: int
    high: int
    low: int
    close: int
    volume: int | None


def _read_bytes(path: Path) -> bytes:
    if path.suffix.lower() != ".zip":
        return path.read_bytes()
    with zipfile.ZipFile(path) as archive:
        members = [member for member in archive.infolist() if not member.is_dir()]
        if len(members) != 1:
            raise ValueError(
                f"{path}: ZIP must contain exactly one data file, found {len(members)}"
            )
        return archive.read(members[0])


def inspect_source_file(path: Path, source_format: SourceFormat) -> SourceInspection:
    if path.suffix.casefold() == ".xlsx":
        return _inspect_xlsx(path, source_format)
    return inspect_delimited_bytes(
        path, _read_bytes(path), source_format.encoding, source_format.delimiter
    )


def resolve_mapping(
    inspection: SourceInspection, configured: str | dict[str, str]
) -> dict[str, str]:
    mapping = inspection.candidate_mapping if configured == "auto" else configured
    if isinstance(mapping, str):
        raise ValueError("source_mapping must be 'auto' or a mapping")
    missing = [
        field
        for field in REQUIRED
        if field not in mapping or mapping[field] not in inspection.headers
    ]
    if missing:
        raise ValueError(
            "cannot resolve required 225Labo columns; "
            f"missing={missing}, candidate_headers={list(inspection.headers)}, "
            f"auto_candidates={inspection.candidate_mapping}. Configure source_mapping explicitly."
        )
    return dict(mapping)


def _parse_int(value: object, field: str, path: Path, row_number: int) -> int:
    try:
        if isinstance(value, bool):
            raise ValueError
        parsed = float(str(value).replace(",", "").strip())
        if not parsed.is_integer():
            raise ValueError
        return int(parsed)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{path}:{row_number}: invalid integer {field}={value!r}") from error


def _parse_date(value: object, path: Path, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip().replace("/", "-"))
    except ValueError as error:
        raise ValueError(f"{path}:{row_number}: invalid date={value!r}") from error


def _parse_time(value: object, path: Path, row_number: int) -> time:
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    try:
        return time.fromisoformat(str(value).strip())
    except ValueError as error:
        raise ValueError(f"{path}:{row_number}: invalid time={value!r}") from error


def read_source_records(path: Path, source_format: SourceFormat) -> Iterator[SourceRecord]:
    if path.suffix.casefold() == ".xlsx":
        yield from _read_xlsx_records(path, source_format)
        return
    raw = _read_bytes(path)
    inspection = inspect_delimited_bytes(path, raw, source_format.encoding, source_format.delimiter)
    mapping = resolve_mapping(inspection, source_format.source_mapping)
    source_hash = hashlib.sha256(raw).hexdigest()
    reader = csv.DictReader(
        io.StringIO(raw.decode(inspection.encoding)), delimiter=inspection.delimiter
    )
    for row_number, row in enumerate(reader, start=2):
        try:
            record_date = _parse_date(row[mapping["date"]], path, row_number)
            record_time = _parse_time(row[mapping["time"]], path, row_number)
        except (KeyError, ValueError) as error:
            raise ValueError(f"{path}:{row_number}: invalid required date/time field") from error
        volume_value = row.get(mapping.get("volume", ""), "")
        yield SourceRecord(
            source_file=str(path),
            source_hash=source_hash,
            source_row_number=row_number,
            trade_date=record_date,
            source_time=record_time,
            open=_parse_int(row[mapping["open"]], "open", path, row_number),
            high=_parse_int(row[mapping["high"]], "high", path, row_number),
            low=_parse_int(row[mapping["low"]], "low", path, row_number),
            close=_parse_int(row[mapping["close"]], "close", path, row_number),
            volume=None
            if volume_value in {None, ""}
            else _parse_int(volume_value, "volume", path, row_number),
        )


def read_source_trade_dates(path: Path, source_format: SourceFormat) -> Iterator[date]:
    """Stream source trade dates without retaining or converting OHLC values."""
    if path.suffix.casefold() == ".xlsx":
        yield from _read_xlsx_trade_dates(path, source_format)
        return
    raw = _read_bytes(path)
    inspection = inspect_delimited_bytes(path, raw, source_format.encoding, source_format.delimiter)
    mapping = resolve_mapping(inspection, source_format.source_mapping)
    reader = csv.DictReader(
        io.StringIO(raw.decode(inspection.encoding)), delimiter=inspection.delimiter
    )
    for row_number, row in enumerate(reader, start=2):
        try:
            yield _parse_date(row[mapping["date"]], path, row_number)
        except (KeyError, ValueError) as error:
            raise ValueError(f"{path}:{row_number}: invalid required date field") from error


def _inspect_xlsx(path: Path, source_format: SourceFormat) -> SourceInspection:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if source_format.worksheet not in workbook.sheetnames:
            raise ValueError(
                f"{path}: worksheet {source_format.worksheet!r} not found; "
                f"available={workbook.sheetnames}"
            )
        worksheet = workbook[source_format.worksheet]
        first = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    finally:
        workbook.close()
    if first is None:
        raise ValueError(f"source workbook {path} is empty")
    headers = tuple(str(value).strip() if value is not None else "" for value in first)
    return inspection_from_headers(path, "xlsx", "worksheet", headers)


def _read_xlsx_records(path: Path, source_format: SourceFormat) -> Iterator[SourceRecord]:
    import openpyxl

    raw = path.read_bytes()
    source_hash = hashlib.sha256(raw).hexdigest()
    inspection = _inspect_xlsx(path, source_format)
    mapping = resolve_mapping(inspection, source_format.source_mapping)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[source_format.worksheet]
        headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        index = {
            str(value).strip() if value is not None else "": position
            for position, value in enumerate(headers)
        }
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            row = {
                header: values[position] if position < len(values) else None
                for header, position in index.items()
            }
            try:
                record_date = _parse_date(row[mapping["date"]], path, row_number)
                record_time = _parse_time(row[mapping["time"]], path, row_number)
            except (KeyError, ValueError) as error:
                raise ValueError(
                    f"{path}:{row_number}: invalid required date/time field"
                ) from error
            volume_value = row.get(mapping.get("volume", ""))
            yield SourceRecord(
                source_file=str(path),
                source_hash=source_hash,
                source_row_number=row_number,
                trade_date=record_date,
                source_time=record_time,
                open=_parse_int(row[mapping["open"]], "open", path, row_number),
                high=_parse_int(row[mapping["high"]], "high", path, row_number),
                low=_parse_int(row[mapping["low"]], "low", path, row_number),
                close=_parse_int(row[mapping["close"]], "close", path, row_number),
                volume=None
                if volume_value in {None, ""}
                else _parse_int(volume_value, "volume", path, row_number),
            )
    finally:
        workbook.close()


def _read_xlsx_trade_dates(path: Path, source_format: SourceFormat) -> Iterator[date]:
    import openpyxl

    inspection = _inspect_xlsx(path, source_format)
    mapping = resolve_mapping(inspection, source_format.source_mapping)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[source_format.worksheet]
        headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        index = {
            str(value).strip() if value is not None else "": position
            for position, value in enumerate(headers)
        }
        date_index = index[mapping["date"]]
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            value = values[date_index] if date_index < len(values) else None
            yield _parse_date(value, path, row_number)
    finally:
        workbook.close()


def write_bronze(
    records: list[SourceRecord],
    output: Path,
    compression: Literal["lz4", "uncompressed", "snappy", "gzip", "brotli", "zstd"] = "zstd",
) -> None:
    """Persist source-shaped, typed records; caller selects a non-raw destination."""
    output.parent.mkdir(parents=True, exist_ok=True)
    pl.DataFrame(
        {
            "source_file": [item.source_file for item in records],
            "source_hash": [item.source_hash for item in records],
            "source_row_number": [item.source_row_number for item in records],
            "source_trade_date": [item.trade_date for item in records],
            "source_time": [item.source_time for item in records],
            "open": [item.open for item in records],
            "high": [item.high for item in records],
            "low": [item.low for item in records],
            "close": [item.close for item in records],
            "volume": [item.volume for item in records],
        }
    ).write_parquet(output, compression=compression)
